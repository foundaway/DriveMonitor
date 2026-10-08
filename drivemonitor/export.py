"""Exportación de tablas y tendencias a CSV o Excel."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable, Sequence


def export_table(path: str | Path, headers: Sequence[str], rows: Iterable[Sequence], sheet: str = "Datos") -> int:
    """Escribe .csv (UTF-8 con BOM, para que Excel respete acentos) o .xlsx según
    la extensión. Devuelve el número de filas."""
    path = Path(path)
    n = 0
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        wb = Workbook(write_only=True)
        ws = wb.create_sheet(title=sheet[:31] or "Datos")
        # En modo write_only el ancho se fija antes de escribir filas.
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(60, len(str(h)) + 4))
        ws.append(list(headers))
        for r in rows:
            ws.append(list(r))
            n += 1
        wb.save(path)
    else:
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(headers)
            for r in rows:
                w.writerow(r)
                n += 1
    return n
